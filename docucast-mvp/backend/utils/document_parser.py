"""Universal document intelligence parser for DocuCast.

Extracts EVERYTHING a document contains — not just body text:

  * PDF   -> text, tables (pdfplumber), embedded images (pypdf), figure /
             graph captions, handwritten-note detection (OCR / vision)
  * PPTX  -> slide text, tables, native chart data (series + categories!),
             images, speaker notes
  * MD    -> body text, markdown tables, image alt text, code blocks
  * TXT   -> plain text with encoding detection

Every visual element (table / image / graph / handwriting) is converted into
a natural-language "narration" string so the podcast script generator can
actually TALK about it — this is the big gap in NotebookLM-style tools,
which typically skip visual content entirely.

Optional enrichment layers (all gracefully skipped when unavailable):
  * pytesseract OCR       (reads printed + handwritten text inside images)
  * Vision LLM captioning (Gemini / OpenRouter / Ollama-llava) describes
    what each image shows, transcribes handwriting, reads charts
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Optional

from pypdf import PdfReader

from backend.utils.vision import describe_image, ocr_image

# ---------------------------------------------------------------------------
# Limits
# ---------------------------------------------------------------------------
MAX_PAGES = 25
MAX_TEXT_CHARS = 14000
MAX_TABLES = 12
MAX_IMAGES = 10
MAX_TABLE_CELL_CHARS = 120
PAGES_LIMIT_NOTICE = "[First 25 pages processed]"
TRUNCATION_NOTICE = "[Content truncated to fit processing budget]"

SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".docx", ".md", ".markdown", ".txt", ".text"}

FIGURE_CAPTION_RE = re.compile(
    r"(?:^|\n)\s*((?:Figure|Fig\.?|Chart|Graph|Diagram|Plot|Exhibit)\s*\.?\s*\d+[.:)\-]?\s*[^\n]{0,220})",
    re.IGNORECASE,
)
TABLE_CAPTION_RE = re.compile(
    r"(?:^|\n)\s*(Table\s*\.?\s*\d+[.:)\-]?\s*[^\n]{0,220})", re.IGNORECASE
)
HANDWRITING_HINT_RE = re.compile(
    r"hand\s*-?\s*written|handwriting|annotat|scribbl|margin note|jotted",
    re.IGNORECASE,
)


@dataclass
class ParsedDocument:
    """Structured result of parsing any supported document."""

    doc_type: str = "unknown"
    text: str = ""
    tables: list = field(default_factory=list)     # {page, title, markdown, narration}
    images: list = field(default_factory=list)     # {page, kind, description, ocr_text, ...}
    figures: list = field(default_factory=list)    # {page, caption}
    charts: list = field(default_factory=list)     # {page, chart_type, narration}
    handwritten_notes: list = field(default_factory=list)  # {page, text, source}
    speaker_notes: list = field(default_factory=list)      # {page, text}
    stats: dict = field(default_factory=dict)
    warnings: list = field(default_factory=list)

    def enriched_text(self) -> str:
        """Merge everything into one narratable brief for the LLM."""
        sections: list[str] = []
        if self.text.strip():
            sections.append("=== DOCUMENT BODY ===\n" + self.text.strip())

        if self.tables:
            lines = ["=== TABLES (data extracted from tables in the document) ==="]
            for t in self.tables:
                lines.append(f"- {t.get('narration', '')}")
            sections.append("\n".join(lines))

        if self.charts:
            lines = ["=== GRAPHS & CHARTS (actual data series read from the charts) ==="]
            for c in self.charts:
                lines.append(f"- {c.get('narration', '')}")
            sections.append("\n".join(lines))

        if self.figures:
            lines = ["=== FIGURE / GRAPH CAPTIONS found in the document ==="]
            for f in self.figures:
                lines.append(f"- (page {f.get('page', '?')}) {f.get('caption', '')}")
            sections.append("\n".join(lines))

        img_lines = []
        for img in self.images:
            desc = img.get("description") or ""
            ocr = img.get("ocr_text") or ""
            bits = []
            if desc:
                bits.append(f"shows: {desc}")
            if ocr:
                bits.append(f"text inside the image: \"{ocr[:400]}\"")
            if bits:
                img_lines.append(
                    f"- Image on page {img.get('page', '?')} ({img.get('kind', 'image')}): "
                    + " | ".join(bits)
                )
        if img_lines:
            sections.append(
                "=== IMAGES (visual content described / read) ===\n" + "\n".join(img_lines)
            )

        if self.handwritten_notes:
            lines = ["=== HANDWRITTEN NOTES / ANNOTATIONS detected ==="]
            for n in self.handwritten_notes:
                lines.append(f"- (page {n.get('page', '?')}) {n.get('text', '')}")
            sections.append("\n".join(lines))

        if self.speaker_notes:
            lines = ["=== PRESENTER / SPEAKER NOTES (hidden notes from the slides) ==="]
            for n in self.speaker_notes:
                lines.append(f"- (slide {n.get('page', '?')}) {n.get('text', '')}")
            sections.append("\n".join(lines))

        return "\n\n".join(sections).strip()

    def summary_payload(self) -> dict:
        """JSON-safe analysis payload for the frontend intelligence panel."""
        return {
            "doc_type": self.doc_type,
            "tables": [
                {k: v for k, v in t.items() if k != "_raw"} for t in self.tables
            ],
            "images": [
                {
                    "page": i.get("page"),
                    "kind": i.get("kind"),
                    "description": i.get("description"),
                    "ocr_text": (i.get("ocr_text") or "")[:600] or None,
                    "width": i.get("width"),
                    "height": i.get("height"),
                }
                for i in self.images
            ],
            "figures": self.figures,
            "charts": self.charts,
            "handwritten_notes": self.handwritten_notes,
            "speaker_notes": self.speaker_notes,
            "stats": self.stats,
            "warnings": self.warnings,
        }


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
def parse_document(file_bytes: bytes, filename: str) -> ParsedDocument:
    """Parse a document of any supported type into a ParsedDocument.

    Raises:
        ValueError: unsupported type or unreadable file.
    """
    name = (filename or "").lower()
    ext = "." + name.rsplit(".", 1)[-1] if "." in name else ""

    if ext == ".pdf":
        return _parse_pdf(file_bytes)
    if ext == ".pptx":
        return _parse_pptx(file_bytes)
    if ext == ".docx":
        return _parse_docx(file_bytes)
    if ext == ".ppt":
        raise ValueError(
            "Legacy .ppt files aren't supported — please re-save the deck as .pptx "
            "(PowerPoint: File → Save As → .pptx) and upload again."
        )
    if ext == ".doc":
        raise ValueError(
            "Legacy .doc files aren't supported — please re-save the document as .docx "
            "(Word: File → Save As → .docx) and upload again."
        )
    if ext in {".md", ".markdown"}:
        return _parse_markdown(file_bytes)
    if ext in {".txt", ".text"}:
        return _parse_text(file_bytes)

    raise ValueError(
        "Unsupported file type. Upload a .pdf, .pptx, .docx, .md or .txt file."
    )


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _parse_pdf(file_bytes: bytes) -> ParsedDocument:
    doc = ParsedDocument(doc_type="pdf")

    # --- pass 1: pdfplumber for text + tables --------------------------------
    plumber_pages = []
    flow_pages = []  # stream-order text: cleaner for side-by-side captions in
    #                  multi-column research papers (layout pass interleaves them)
    try:
        import pdfplumber

        with pdfplumber.open(BytesIO(file_bytes)) as pdf:
            total_pages = len(pdf.pages)
            for idx, page in enumerate(pdf.pages[:MAX_PAGES]):
                page_no = idx + 1
                try:
                    page_text = page.extract_text() or ""
                except Exception:
                    page_text = ""
                plumber_pages.append(page_text)
                try:
                    flow_pages.append(page.extract_text(use_text_flow=True) or "")
                except Exception:
                    flow_pages.append("")

                # Tables
                if len(doc.tables) < MAX_TABLES:
                    try:
                        raw_tables = page.extract_tables() or []
                    except Exception:
                        raw_tables = []
                    if not raw_tables:
                        # Borderless tables: retry with text alignment strategy,
                        # accepting only results that genuinely look tabular.
                        try:
                            candidates = page.extract_tables(
                                {"vertical_strategy": "text", "horizontal_strategy": "text"}
                            ) or []
                            raw_tables = [t for t in candidates if _looks_tabular(t)]
                        except Exception:
                            raw_tables = []
                    for raw in raw_tables:
                        if len(doc.tables) >= MAX_TABLES:
                            break
                        table = _table_to_struct(raw, page_no)
                        if table:
                            doc.tables.append(table)
    except Exception as exc:
        doc.warnings.append(f"Advanced PDF layout analysis unavailable ({exc}); using basic extraction.")
        total_pages = 0

    # --- pass 2: pypdf fallback text + encrypted handling + images -----------
    try:
        reader = PdfReader(BytesIO(file_bytes))
    except Exception as exc:
        if plumber_pages:
            reader = None
            total_pages = total_pages or len(plumber_pages)
        else:
            raise ValueError("Unable to read PDF. The file may be corrupted.") from exc

    if reader is not None:
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError(
                    "This PDF is password-protected. Remove the password and upload again."
                )
        total_pages = len(reader.pages)

        pages_to_read = min(total_pages, MAX_PAGES)
        for i in range(pages_to_read):
            # Fill in text where pdfplumber came up empty
            if i >= len(plumber_pages) or not plumber_pages[i].strip():
                try:
                    fallback = reader.pages[i].extract_text() or ""
                except Exception:
                    fallback = ""
                if i < len(plumber_pages):
                    plumber_pages[i] = fallback
                else:
                    plumber_pages.append(fallback)

        # Embedded images
        image_budget = MAX_IMAGES
        for i in range(pages_to_read):
            if image_budget <= 0:
                break
            try:
                page_images = reader.pages[i].images
            except Exception:
                page_images = []
            for img_file in page_images:
                if image_budget <= 0:
                    break
                try:
                    entry = _analyze_image(img_file.data, page=i + 1,
                                           page_text=plumber_pages[i] if i < len(plumber_pages) else "")
                except Exception:
                    entry = None
                if entry:
                    doc.images.append(entry)
                    image_budget -= 1

    text = clean_text("\n".join(plumber_pages))

    # Figure / graph captions: prefer the stream-order pass (side-by-side
    # captions stay intact), fall back to the layout pass per page.
    caption_pages = [
        flow_pages[i] if i < len(flow_pages) and flow_pages[i].strip() else page_text
        for i, page_text in enumerate(plumber_pages)
    ]
    doc.figures = _extract_figure_captions(caption_pages)

    # Handwriting: any image whose OCR/vision result looked handwritten
    for img in doc.images:
        if img.get("handwritten") and (img.get("ocr_text") or img.get("description")):
            doc.handwritten_notes.append({
                "page": img.get("page"),
                "text": img.get("ocr_text") or img.get("description"),
                "source": "image analysis",
            })

    if total_pages > MAX_PAGES:
        text = f"{text}\n\n{PAGES_LIMIT_NOTICE}"
        doc.warnings.append(f"Document has {total_pages} pages; the first {MAX_PAGES} were processed.")

    doc.text = _truncate(text)
    doc.stats = {
        "pages": total_pages,
        "pages_processed": min(total_pages, MAX_PAGES) if total_pages else len(plumber_pages),
        "words": len(doc.text.split()),
        "tables": len(doc.tables),
        "images": len(doc.images),
        "figures": len(doc.figures),
        "handwritten_notes": len(doc.handwritten_notes),
    }

    if not doc.text.strip() and not doc.tables and not doc.images:
        raise ValueError(
            "No readable content found in this PDF. It may be a scanned/image-only "
            "document — enable OCR (install tesseract) or set a vision-capable API key "
            "(GEMINI_API_KEY) so DocuCast can read scanned pages."
        )
    if not doc.text.strip() and (doc.tables or doc.images):
        # Scanned doc but we got visual intelligence — still narratable.
        doc.warnings.append(
            "No selectable body text found — the podcast is built from tables/images that could be analyzed."
        )

    return doc


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------
def _parse_pptx(file_bytes: bytes) -> ParsedDocument:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise ValueError("PPTX support requires python-pptx on the server.") from exc

    doc = ParsedDocument(doc_type="pptx")
    try:
        prs = Presentation(BytesIO(file_bytes))
    except Exception as exc:
        raise ValueError("Unable to read this PowerPoint file. It may be corrupted or a legacy .ppt.") from exc

    slide_texts: list[str] = []
    image_budget = MAX_IMAGES

    slides = list(prs.slides)
    for idx, slide in enumerate(slides[:MAX_PAGES]):
        slide_no = idx + 1
        chunks: list[str] = []
        title = ""
        try:
            if slide.shapes.title and slide.shapes.title.text:
                title = slide.shapes.title.text.strip()
        except Exception:
            pass

        for shape in slide.shapes:
            # Text
            try:
                if shape.has_text_frame and shape.text.strip():
                    if not title or shape.text.strip() != title:
                        chunks.append(shape.text.strip())
            except Exception:
                pass

            # Tables
            try:
                if shape.has_table and len(doc.tables) < MAX_TABLES:
                    raw = [
                        [cell.text for cell in row.cells]
                        for row in shape.table.rows
                    ]
                    table = _table_to_struct(raw, slide_no)
                    if table:
                        doc.tables.append(table)
            except Exception:
                pass

            # Native charts — we can read the actual data series!
            try:
                if shape.has_chart:
                    chart_entry = _pptx_chart_to_struct(shape.chart, slide_no)
                    if chart_entry:
                        doc.charts.append(chart_entry)
            except Exception:
                pass

            # Pictures
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and image_budget > 0:
                    blob = shape.image.blob
                    entry = _analyze_image(blob, page=slide_no, page_text=" ".join(chunks))
                    if entry:
                        doc.images.append(entry)
                        image_budget -= 1
            except Exception:
                pass

        header = f"[Slide {slide_no}{': ' + title if title else ''}]"
        body = "\n".join(chunks)
        slide_texts.append(f"{header}\n{body}" if body or title else header)

        # Speaker notes — hidden context most tools throw away
        try:
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    doc.speaker_notes.append({"page": slide_no, "text": notes[:800]})
        except Exception:
            pass

    for img in doc.images:
        if img.get("handwritten") and (img.get("ocr_text") or img.get("description")):
            doc.handwritten_notes.append({
                "page": img.get("page"),
                "text": img.get("ocr_text") or img.get("description"),
                "source": "slide image analysis",
            })

    text = clean_text("\n\n".join(slide_texts))
    if len(slides) > MAX_PAGES:
        doc.warnings.append(f"Deck has {len(slides)} slides; the first {MAX_PAGES} were processed.")

    doc.text = _truncate(text)
    doc.figures = _extract_figure_captions(slide_texts)
    doc.stats = {
        "pages": len(slides),
        "pages_processed": min(len(slides), MAX_PAGES),
        "words": len(doc.text.split()),
        "tables": len(doc.tables),
        "images": len(doc.images),
        "charts": len(doc.charts),
        "speaker_notes": len(doc.speaker_notes),
        "handwritten_notes": len(doc.handwritten_notes),
    }

    if not doc.text.strip() and not doc.tables and not doc.charts and not doc.images:
        raise ValueError("This presentation appears to be empty — no text, tables, charts or images found.")
    return doc


def _pptx_chart_to_struct(chart, slide_no: int) -> Optional[dict]:
    """Read real data out of a native PPTX chart and narrate it."""
    try:
        chart_type = str(chart.chart_type).split(".")[-1].split(" ")[0].replace("_", " ").title()
    except Exception:
        chart_type = "Chart"
    try:
        title = chart.chart_title.text_frame.text if chart.has_title else ""
    except Exception:
        title = ""

    series_bits = []
    categories = []
    try:
        plot = chart.plots[0]
        categories = [str(c) for c in list(plot.categories)[:8]]
        for s in list(plot.series)[:4]:
            vals = [v for v in list(s.values)[:8]]
            pretty = ", ".join(_fmt_num(v) for v in vals)
            series_bits.append(f"series '{s.name}': {pretty}")
    except Exception:
        pass

    narration = f"{chart_type} on slide {slide_no}"
    if title:
        narration += f" titled '{title}'"
    if categories:
        narration += f"; categories: {', '.join(categories)}"
    if series_bits:
        narration += f"; {'; '.join(series_bits)}"
        # Add trend insight
        trend = _trend_insight(series_bits)
        if trend:
            narration += f". {trend}"
    if not categories and not series_bits and not title:
        narration += " (data could not be read)"

    return {"page": slide_no, "chart_type": chart_type, "title": title or None, "narration": narration}


def _fmt_num(v) -> str:
    if v is None:
        return "n/a"
    try:
        f = float(v)
        if f.is_integer():
            return str(int(f))
        return f"{f:,.2f}".rstrip("0").rstrip(".")
    except Exception:
        return str(v)


def _trend_insight(series_bits: list) -> str:
    """Tiny heuristic: mention rising/falling trend of the first series."""
    try:
        nums = re.findall(r"-?\d+(?:\.\d+)?", series_bits[0].split(":", 1)[1])
        vals = [float(n) for n in nums]
        if len(vals) >= 3:
            if vals[-1] > vals[0] * 1.1:
                return "The trend rises overall from start to end."
            if vals[-1] < vals[0] * 0.9:
                return "The trend falls overall from start to end."
            return "Values stay roughly flat across the range."
    except Exception:
        pass
    return ""


# ---------------------------------------------------------------------------
# DOCX (Word)
# ---------------------------------------------------------------------------
def _parse_docx(file_bytes: bytes) -> ParsedDocument:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise ValueError("DOCX support requires python-docx on the server.") from exc

    doc = ParsedDocument(doc_type="docx")
    try:
        d = docx.Document(BytesIO(file_bytes))
    except Exception as exc:
        raise ValueError(
            "Unable to read this Word document. It may be corrupted or a legacy .doc file."
        ) from exc

    chunks: list[str] = []
    body = d.element.body
    image_budget = MAX_IMAGES

    # Walk the body in order so headings/paragraphs/tables keep document order.
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]

        if tag == "p":
            para = Paragraph(child, d)
            try:
                style = (para.style.name or "").lower()
            except Exception:
                style = ""
            text = para.text.strip()
            if not text:
                continue
            if style.startswith("heading"):
                try:
                    level = int(style.replace("heading", "").strip() or "1")
                except ValueError:
                    level = 1
                chunks.append("#" * min(level, 4) + " " + text)
            else:
                chunks.append(text)

        elif tag == "tbl":
            table = Table(child, d)
            if len(doc.tables) < MAX_TABLES:
                raw = [[cell.text for cell in row.cells] for row in table.rows]
                t = _table_to_struct(raw, page=1)
                if t:
                    doc.tables.append(t)

    # Embedded images (inline shapes reference image parts)
    try:
        for rel_id, rel in d.part.rels.items():
            if image_budget <= 0:
                break
            if not rel.reltype.endswith("/image"):
                continue
            try:
                entry = _analyze_image(rel.target_part.blob, page=1, page_text="")
            except Exception:
                entry = None
            if entry:
                doc.images.append(entry)
                image_budget -= 1
    except Exception:
        pass

    text = clean_text("\n\n".join(chunks))
    doc.figures = _extract_figure_captions([text])
    doc.text = _truncate(text)

    doc.stats = {
        "pages": 1,
        "words": len(doc.text.split()),
        "tables": len(doc.tables),
        "images": len(doc.images),
        "figures": len(doc.figures),
    }

    if not doc.text.strip() and not doc.tables and not doc.images:
        raise ValueError("This Word document appears to be empty — no text, tables or images found.")
    return doc


# ---------------------------------------------------------------------------
# Markdown / plain text
# ---------------------------------------------------------------------------
MD_TABLE_RE = re.compile(
    r"((?:^\|.+\|\s*$\n)+)", re.MULTILINE
)
MD_IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)\s]+)[^)]*\)")
MD_CODEBLOCK_RE = re.compile(r"```(\w+)?\n(.*?)```", re.DOTALL)


def _decode(file_bytes: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return file_bytes.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return file_bytes.decode("utf-8", errors="replace")


def _parse_markdown(file_bytes: bytes) -> ParsedDocument:
    doc = ParsedDocument(doc_type="markdown")
    raw = _decode(file_bytes)

    # Markdown tables → structured
    for m in list(MD_TABLE_RE.finditer(raw))[:MAX_TABLES]:
        block = m.group(1)
        rows = []
        for line in block.strip().splitlines():
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                continue  # separator row
            rows.append(cells)
        table = _table_to_struct(rows, page=1)
        if table:
            doc.tables.append(table)

    # Image references (alt text is real content!)
    for m in list(MD_IMAGE_RE.finditer(raw))[:MAX_IMAGES]:
        alt, src = m.group(1).strip(), m.group(2)
        doc.images.append({
            "page": 1,
            "kind": "linked image",
            "description": alt or f"image referenced at {src.split('/')[-1]}",
            "ocr_text": None,
            "handwritten": False,
        })

    # Code blocks — mention them, don't read them char by char
    code_langs = [m.group(1) or "code" for m in MD_CODEBLOCK_RE.finditer(raw)]
    if code_langs:
        doc.warnings.append(
            f"Contains {len(code_langs)} code block(s) ({', '.join(sorted(set(code_langs))[:4])}) — summarized, not read verbatim."
        )

    # Strip markdown syntax for the narratable body
    text = MD_CODEBLOCK_RE.sub(lambda m: f"\n[Code block in {m.group(1) or 'code'} — omitted from narration]\n", raw)
    text = MD_TABLE_RE.sub("\n", text)
    text = MD_IMAGE_RE.sub(lambda m: f" (image: {m.group(1)}) " if m.group(1) else " ", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)          # links → label
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)     # headings
    text = re.sub(r"[*_]{1,3}([^*_\n]+)[*_]{1,3}", r"\1", text)   # bold/italic
    text = re.sub(r"^>\s?", "", text, flags=re.MULTILINE)          # quotes
    text = re.sub(r"`([^`]+)`", r"\1", text)                       # inline code

    doc.text = _truncate(clean_text(text))
    doc.figures = _extract_figure_captions([raw])
    doc.stats = {
        "pages": 1,
        "words": len(doc.text.split()),
        "tables": len(doc.tables),
        "images": len(doc.images),
        "code_blocks": len(code_langs),
    }
    if not doc.text.strip() and not doc.tables:
        raise ValueError("This markdown file appears to be empty.")
    return doc


def _parse_text(file_bytes: bytes) -> ParsedDocument:
    doc = ParsedDocument(doc_type="text")
    raw = _decode(file_bytes)
    doc.text = _truncate(clean_text(raw))
    doc.figures = _extract_figure_captions([raw])
    doc.stats = {"pages": 1, "words": len(doc.text.split())}
    if not doc.text.strip():
        raise ValueError("This text file appears to be empty.")
    return doc


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------
def _looks_tabular(raw_rows) -> bool:
    """Guard against text-strategy false positives (prose parsed as 'tables')."""
    if not raw_rows:
        return False
    rows = [[(c or "").strip() for c in (r or [])] for r in raw_rows]
    rows = [r for r in rows if any(r)]
    if not (3 <= len(rows) <= 60):
        return False
    width = max(len(r) for r in rows)
    if not (2 <= width <= 8):
        return False
    multi_cell = sum(1 for r in rows if sum(1 for c in r if c) >= 2)
    if multi_cell / len(rows) < 0.6:
        return False
    # Prose rows tend to have very long cells; tabular cells are short.
    long_cells = sum(1 for r in rows for c in r if len(c) > 80)
    total_cells = sum(1 for r in rows for c in r if c)
    return total_cells > 0 and long_cells / total_cells < 0.2


def _table_to_struct(raw_rows, page: int) -> Optional[dict]:
    """Convert a raw table (list of rows) into markdown + spoken narration."""
    if not raw_rows:
        return None
    rows = []
    for r in raw_rows:
        cells = [
            re.sub(r"\s+", " ", (c or "")).strip()[:MAX_TABLE_CELL_CHARS]
            for c in (r or [])
        ]
        if any(cells):
            rows.append(cells)
    if len(rows) < 2 or max(len(r) for r in rows) < 2:
        return None  # not a real table

    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, body = rows[0], rows[1:8]  # cap narration at 7 data rows

    md_lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    for r in rows[1:]:
        md_lines.append("| " + " | ".join(r) + " |")

    # Spoken narration: "row label — col: val, col: val"
    spoken_rows = []
    for r in body:
        label = r[0] or "row"
        pairs = [
            f"{header[i] or f'column {i + 1}'}: {r[i]}"
            for i in range(1, width)
            if r[i]
        ]
        if pairs:
            spoken_rows.append(f"{label} — {', '.join(pairs[:5])}")
    narration = (
        f"Table on page {page} with columns [{', '.join(h for h in header if h)}] "
        f"and {len(rows) - 1} rows. Key rows: " + "; ".join(spoken_rows[:5])
        if spoken_rows
        else f"Table on page {page} with {len(rows) - 1} rows and {width} columns."
    )

    return {
        "page": page,
        "columns": header,
        "row_count": len(rows) - 1,
        "markdown": "\n".join(md_lines[:12]),
        "narration": narration,
    }


def _analyze_image(image_bytes: bytes, page: int, page_text: str = "") -> Optional[dict]:
    """Analyze one embedded image: size sanity, OCR, vision description, kind guess."""
    if not image_bytes or len(image_bytes) < 500:
        return None  # skip tiny icons / decorations

    width = height = None
    try:
        from PIL import Image

        with Image.open(io.BytesIO(image_bytes)) as im:
            width, height = im.size
        if width < 64 or height < 64:
            return None  # decorative
    except Exception:
        pass

    # Guess the kind from nearby captions
    kind = "image"
    ctx = (page_text or "")[:1500]
    if re.search(r"\b(chart|graph|plot|histogram|scatter|bar chart|line chart|pie)\b", ctx, re.I):
        kind = "graph/chart"
    elif re.search(r"\b(figure|fig\.|diagram|architecture|pipeline|flow)\b", ctx, re.I):
        kind = "figure/diagram"
    elif HANDWRITING_HINT_RE.search(ctx):
        kind = "handwritten note"

    ocr_text = ocr_image(image_bytes)  # None if OCR unavailable
    vision = describe_image(image_bytes, context_hint=ctx[:400])  # None if no vision provider

    description = None
    handwritten = kind == "handwritten note"
    if vision:
        description = vision.get("description")
        if vision.get("kind"):
            kind = vision["kind"]
        if vision.get("handwritten"):
            handwritten = True
        if vision.get("text") and not ocr_text:
            ocr_text = vision["text"]

    if not description and not ocr_text:
        # Nothing readable — still report presence with the heuristic kind
        description = (
            f"{kind} (~{width}x{height}px)" if width else kind
        ) + " — no OCR/vision provider configured to read its contents"

    return {
        "page": page,
        "kind": kind,
        "width": width,
        "height": height,
        "ocr_text": (ocr_text or "").strip()[:800] or None,
        "description": description,
        "handwritten": handwritten,
    }


def _extract_figure_captions(pages: list) -> list:
    figures = []
    seen = set()
    for idx, page_text in enumerate(pages):
        if not page_text:
            continue
        for m in FIGURE_CAPTION_RE.finditer(page_text):
            cap = re.sub(r"\s+", " ", m.group(1)).strip()
            key = cap.lower()[:60]
            if key not in seen and len(cap) > 8:
                seen.add(key)
                figures.append({"page": idx + 1, "caption": cap[:240]})
        for m in TABLE_CAPTION_RE.finditer(page_text):
            cap = re.sub(r"\s+", " ", m.group(1)).strip()
            key = cap.lower()[:60]
            if key not in seen and len(cap) > 8:
                seen.add(key)
                figures.append({"page": idx + 1, "caption": cap[:240]})
    return figures[:20]


def clean_text(text: str) -> str:
    """Normalize whitespace while preserving paragraph breaks."""
    if not text:
        return ""
    lines = text.splitlines()
    cleaned_lines = []
    for line in lines:
        stripped = " ".join(line.split())
        cleaned_lines.append(stripped)
    cleaned = "\n".join(cleaned_lines)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def _truncate(text: str, limit: int = MAX_TEXT_CHARS) -> str:
    if len(text) <= limit:
        return text
    cutoff = text[:limit]
    last_boundary = -1
    for punctuation in (". ", "! ", "? ", ".\n", "!\n", "?\n"):
        idx = cutoff.rfind(punctuation)
        if idx > last_boundary:
            last_boundary = idx
    if last_boundary == -1:
        last_space = cutoff.rfind(" ")
        out = cutoff[:last_space].strip() if last_space > 0 else cutoff.strip()
    else:
        out = cutoff[: last_boundary + 1].strip()
    return f"{out} {TRUNCATION_NOTICE}"


# ---------------------------------------------------------------------------
# Backwards compatibility
# ---------------------------------------------------------------------------
def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Legacy API: returns enriched text for a PDF."""
    return _parse_pdf(file_bytes).enriched_text()
